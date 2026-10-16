"""Excel and PDF exports.

Excel workbooks are verified two ways: re-read with openpyxl and recalculated
with FinPilot's own engine, and (when LibreOffice is installed) recalculated
headlessly by LibreOffice and compared cell by cell. Microsoft Excel itself is
not used by this test suite.
"""
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from finpilot.core import hotel as H
from finpilot.core import hotel_sheet as hs
from finpilot.curriculum import progression as prog
from finpilot.export import excel, pdf


@pytest.fixture
def model_xlsx(tmp_path):
    sheet = hs.build_sheet(H.default_forecast(), H.default_budget(), include_calcs=True)
    path = excel.export_hotel_model(tmp_path / "model.xlsx", sheet, H.default_forecast(), H.default_budget(),
                                    "Test model", "Protect winter rate.")
    return sheet, path


def test_workbook_structure_and_real_formulas(model_xlsx):
    sheet, path = model_xlsx
    wb = load_workbook(path)
    assert {"Cover", "Model", "Scenario inputs", "Scenarios", "Bridge", "Checks & notes",
            "Recommendation"} <= set(wb.sheetnames)
    ws = wb["Model"]
    assert ws[f"B{hs.R_AVAIL}"].value == "=$B$4*B9"
    assert str(ws[f"N{hs.R_RREV}"].value).startswith("=SUM(")
    assert wb.calculation.fullCalcOnLoad
    assert "fictional" in " ".join(str(c.value) for row in wb["Cover"].iter_rows() for c in row if c.value).lower()


def test_exported_formulas_recalculate_to_the_same_values(model_xlsx):
    sheet, path = model_xlsx
    back = excel.read_back_sheet(path, "Model")
    compared = 0
    for (r, c), cell in sheet.cells.items():
        if cell.is_formula:
            a, b = sheet.value((r, c)), back.value((r, c))
            if isinstance(a, (int, float)) and not isinstance(a, bool):
                assert b == pytest.approx(a, rel=1e-12, abs=1e-9), (r, c)
            else:
                assert a == b, (r, c)
            compared += 1
    assert compared > 150


def test_lab_sheet_export(tmp_path):
    sheet = hs.build_sheet(H.default_forecast(), H.default_budget(), include_calcs=False)
    p = excel.export_lab_sheet(tmp_path / "lab.xlsx", sheet, "Lab")
    assert load_workbook(p)["Model"]["B4"].value == 180


def _soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


@pytest.mark.skipif(_soffice() is None, reason="LibreOffice not installed")
def test_libreoffice_recalculation_matches_finpilot(model_xlsx, tmp_path):
    sheet, path = model_xlsx
    out = tmp_path / "lo"
    out.mkdir()
    subprocess.run([_soffice(), "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(out / "model.xlsx", data_only=True)
    ws = wb["Model"]
    checked = 0
    for (r, c), cell in sheet.cells.items():
        if not cell.is_formula:
            continue
        mine = sheet.value((r, c))
        lo = ws.cell(row=r + 1, column=c + 1).value
        if isinstance(mine, bool):
            assert lo in (mine, int(mine)), (r, c, lo)
        elif isinstance(mine, (int, float)):
            assert lo == pytest.approx(mine, rel=1e-9, abs=1e-6), (r, c, lo)
        else:
            assert lo == mine, (r, c, lo)
        checked += 1
    assert checked > 150
    assert round(ws[f"N{hs.R_RREV}"].value) == 17_484_066
    sc = wb["Scenarios"]
    values = [c.value for row in sc.iter_rows() for c in row if isinstance(c.value, (int, float))]
    assert any(round(v) == 17_366_058 for v in values) and any(round(v) == 17_484_066 for v in values)
    br = wb["Bridge"]
    flags = [c.value for row in br.iter_rows() for c in row if isinstance(c.value, bool)]
    assert flags and all(flags)


def test_pdf_reports(tmp_path, db, cur):
    p1 = pdf.project_report_pdf(tmp_path / "project.pdf", "Test", H.default_forecast(), H.default_budget(),
                                "Protect winter rate.", [("Rooms sold", "pass", "ok")])
    p2 = pdf.competency_report_pdf(tmp_path / "comp.pdf", prog.competency_report(cur, db))
    for p in (p1, p2):
        data = Path(p).read_bytes()
        assert data.startswith(b"%PDF") and len(data) > 5000
