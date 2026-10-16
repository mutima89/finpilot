"""Excel workbook export with live formulas.

FinPilot writes real Excel formulas (not pasted values) for every calculated
cell. openpyxl cannot calculate formulas, so the workbook is flagged to fully
recalculate when it is opened (``fullCalcOnLoad``). FinPilot does NOT run
Excel; values shown before recalculation depend on the spreadsheet program.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties

from .. import __app_name__, __tagline__, __version__
from ..core.hotel import MONTHS, HotelAssumptions, validate
from ..core.sheet import Sheet

NAVY = "101D2B"
TEAL = "59D6C0"
AMBER = "F3B65B"
CREAM = "F5F1E8"
SLATE = "5B6675"
INPUT_BLUE = "1F4FBF"

# Functions newer than Excel 2010 must be written with the _xlfn. prefix in the file.
_XLFN = {"XLOOKUP", "IFS", "IFNA", "CONCAT", "DAYS", "MAXIFS", "MINIFS"}
_FN_RE = re.compile(r"\b(" + "|".join(sorted(_XLFN, key=len, reverse=True)) + r")\(", re.I)

NUMBER_FORMATS = {
    "currency": '"$"#,##0.00;("$"#,##0.00);"-"',
    "currency0": '"$"#,##0;("$"#,##0);"-"',
    "number": '#,##0.00;(#,##0.00);"-"',
    "number0": '#,##0;(#,##0);"-"',
    "percent": '0.0%;-0.0%;"-"',
    "percent0": '0%;-0%;"-"',
    "date": "dd-mmm-yyyy",
    "general": "General",
    "text": "@",
}


def excel_formula(raw: str) -> str:
    """Convert a FinPilot formula to the form stored in .xlsx files."""
    out = []
    # only prefix outside string literals
    parts = re.split(r'("(?:[^"]|"")*")', raw)
    for p in parts:
        if p.startswith('"'):
            out.append(p)
        else:
            out.append(_FN_RE.sub(lambda m: "_xlfn." + m.group(1).upper() + "(", p))
    return "".join(out)


def _style_cell(c, role: str, fmt: str, sheet_cell_is_formula: bool) -> None:
    c.number_format = NUMBER_FORMATS.get(fmt, "General")
    if role == "input":
        c.font = Font(name="Calibri", color=INPUT_BLUE)
        c.fill = PatternFill("solid", fgColor="FFF7E0")
    elif role == "output":
        c.font = Font(name="Calibri", bold=True, color="000000")
        c.fill = PatternFill("solid", fgColor="E6F7F3")
    elif role == "check":
        c.font = Font(name="Calibri", italic=True, color="1B7F5A")
    elif role == "label":
        c.font = Font(name="Calibri", color="222222")
    else:
        c.font = Font(name="Calibri", color="000000")


def write_sheet(ws, sheet: Sheet) -> None:
    """Write a FinPilot sheet to an openpyxl worksheet, keeping formulas."""
    for (r, col), cell in sorted(sheet.cells.items()):
        if not cell.raw and not cell.role:
            continue
        xc = ws.cell(row=r + 1, column=col + 1)
        if cell.is_formula:
            xc.value = excel_formula(cell.raw)
        else:
            v = cell.literal
            xc.value = v
        _style_cell(xc, cell.role, cell.fmt, cell.is_formula)
        if cell.note:
            from openpyxl.comments import Comment
            xc.comment = Comment(cell.note, __app_name__)
    ws.column_dimensions["A"].width = 38
    for col in range(2, sheet.cols + 2):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "B8" if ws.title == "Model" else None


def _cover(wb: Workbook, title: str, subtitle: str, notes: list[str]) -> None:
    ws = wb.active
    ws.title = "Cover"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    for r in range(1, 6):
        for c in range(1, 4):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=NAVY)
    ws["B2"] = __app_name__
    ws["B2"].font = Font(name="Calibri", size=26, bold=True, color=TEAL)
    ws["B3"] = __tagline__
    ws["B3"].font = Font(name="Calibri", size=12, italic=True, color=AMBER)
    ws["B4"] = "Interactive Financial Modeling Academy"
    ws["B4"].font = Font(name="Calibri", size=10, color="FFFFFF")
    ws["B7"] = title
    ws["B7"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
    ws["B8"] = subtitle
    ws["B8"].font = Font(name="Calibri", size=11, color=SLATE)
    ws["B9"] = f"Exported {dt.datetime.now():%d %B %Y %H:%M} with {__app_name__} {__version__}"
    ws["B9"].font = Font(name="Calibri", size=9, color=SLATE)
    row = 11
    for n in notes:
        ws.cell(row=row, column=2, value=n).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=2).font = Font(name="Calibri", size=10, color="222222")
        row += 1
    row += 1
    ws.cell(row=row, column=2, value="Legend").font = Font(bold=True, color=NAVY)
    legend = [("Blue text on light amber: input / assumption", "input"),
              ("Black text: calculation (formula)", "calc"),
              ("Bold on light teal: key output", "output"),
              ("Green italic: model check (TRUE = passes)", "check")]
    for text, role in legend:
        row += 1
        c = ws.cell(row=row, column=2, value=text)
        _style_cell(c, role, "general", False)


def export_hotel_model(path: Path, sheet: Sheet, forecast: HotelAssumptions, budget: HotelAssumptions,
                       project_name: str = "Hotel revenue model", recommendation: str = "") -> Path:
    path = Path(path)
    wb = Workbook()
    _cover(wb, project_name, f"{forecast.hotel_name} - 12-month revenue model, {forecast.year}, {forecast.currency}",
           ["All data in this workbook is fictional and for training purposes only.",
            "Sheet 'Model' contains the live model built in FinPilot: inputs, formulas, budget comparison and checks. "
            "Every calculated cell is an Excel formula.",
            "Sheet 'Scenarios' compares the budget and forecast scenarios using formulas that reference sheet "
            "'Scenario inputs'.",
            "Sheet 'Bridge' splits the room revenue variance into volume and rate effects (formulas).",
            "Calculation: the workbook is set to recalculate fully when opened. FinPilot did not open or "
            "recalculate this file in Microsoft Excel; if values look blank in a viewer that does not calculate, "
            "open it in Excel and press F9.",
            "Methodology: Occupancy = Rooms Sold / Available Room Nights; ADR = Room Revenue / Rooms Sold; "
            "RevPAR = Room Revenue / Available Room Nights. Annual ratios are weighted from totals. Rooms sold "
            "are not rounded so that RevPAR = ADR x Occupancy holds exactly.",
            "Simplifications: F&B revenue = rooms sold x spend per occupied room; other revenue = % of room "
            "revenue. Real hotels forecast these by outlet and segment."])
    ws = wb.create_sheet("Model")
    write_sheet(ws, sheet)
    for name, target in sheet.names.items():
        parts = target.split(":")
        ref = ":".join("$" + re.sub(r"(\d+)", r"$\1", p) for p in parts)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Model!{ref}")

    _scenario_sheets(wb, forecast, budget)
    _bridge_sheet(wb)
    _checks_sheet(wb, forecast)
    if recommendation.strip():
        rs = wb.create_sheet("Recommendation")
        rs.column_dimensions["A"].width = 110
        rs["A1"] = "Management recommendation"
        rs["A1"].font = Font(bold=True, size=14, color=NAVY)
        rs["A3"] = recommendation
        rs["A3"].alignment = Alignment(wrap_text=True, vertical="top")
        rs.row_dimensions[3].height = 220
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _scenario_sheets(wb: Workbook, forecast: HotelAssumptions, budget: HotelAssumptions) -> None:
    si = wb.create_sheet("Scenario inputs")
    si["A1"] = "Scenario inputs (fictional)"
    si["A1"].font = Font(bold=True, size=13, color=NAVY)
    si["A2"] = "Rooms"
    si["B2"] = forecast.rooms
    si["A3"] = "Year"
    si["B3"] = forecast.year
    for c in ("B2", "B3"):
        _style_cell(si[c], "input", "number0" if c == "B2" else "general", False)
    headers = ["Month"] + MONTHS
    blocks = [("Budget occupancy", budget, "occupancy", "percent"), ("Budget ADR", budget, "adr", "currency"),
              ("Forecast occupancy", forecast, "occupancy", "percent"), ("Forecast ADR", forecast, "adr", "currency")]
    for j, h in enumerate(headers):
        si.cell(row=5, column=1 + j, value=h).font = Font(bold=True)
    for i, (label, a, attr, fmt) in enumerate(blocks):
        r = 6 + i
        si.cell(row=r, column=1, value=label)
        for m in range(12):
            c = si.cell(row=r, column=2 + m, value=float(getattr(a.months[m], attr)))
            _style_cell(c, "input", fmt, False)
    si.column_dimensions["A"].width = 24

    sc = wb.create_sheet("Scenarios")
    sc["A1"] = "Budget versus forecast - room revenue (formulas)"
    sc["A1"].font = Font(bold=True, size=13, color=NAVY)
    for j, h in enumerate(["Line"] + MONTHS + ["Full year"]):
        sc.cell(row=3, column=1 + j, value=h).font = Font(bold=True)
    rows = [
        (4, "Days", lambda col: f"=DAY(EOMONTH(DATE('Scenario inputs'!$B$3,{col - 1},1),0))", "number0"),
        (5, "Available room nights", lambda col: f"='Scenario inputs'!$B$2*{get_column_letter(col)}4", "number0"),
        (6, "Budget rooms sold", lambda col: f"={get_column_letter(col)}5*'Scenario inputs'!{get_column_letter(col)}6", "number0"),
        (7, "Budget room revenue", lambda col: f"={get_column_letter(col)}6*'Scenario inputs'!{get_column_letter(col)}7", "currency0"),
        (8, "Forecast rooms sold", lambda col: f"={get_column_letter(col)}5*'Scenario inputs'!{get_column_letter(col)}8", "number0"),
        (9, "Forecast room revenue", lambda col: f"={get_column_letter(col)}8*'Scenario inputs'!{get_column_letter(col)}9", "currency0"),
        (10, "Variance (forecast - budget)", lambda col: f"={get_column_letter(col)}9-{get_column_letter(col)}7", "currency0"),
        (11, "Budget RevPAR", lambda col: f"=IF({get_column_letter(col)}5=0,NA(),{get_column_letter(col)}7/{get_column_letter(col)}5)", "currency"),
        (12, "Forecast RevPAR", lambda col: f"=IF({get_column_letter(col)}5=0,NA(),{get_column_letter(col)}9/{get_column_letter(col)}5)", "currency"),
    ]
    for r, label, f, fmt in rows:
        sc.cell(row=r, column=1, value=label)
        for col in range(2, 14):
            c = sc.cell(row=r, column=col, value=f(col))
            _style_cell(c, "output" if r in (7, 9, 10) else "calc", fmt, True)
        tot = sc.cell(row=r, column=14)
        if r in (11, 12):
            src = 7 if r == 11 else 9
            tot.value = f"=IF(N5=0,NA(),N{src}/N5)"
        else:
            tot.value = f"=SUM(B{r}:M{r})"
        _style_cell(tot, "output", fmt, True)
    sc.column_dimensions["A"].width = 30
    for col in range(2, 15):
        sc.column_dimensions[get_column_letter(col)].width = 13

    chart = BarChart()
    chart.type = "col"
    chart.title = "Room revenue: budget vs forecast"
    chart.y_axis.title = "USD"
    chart.y_axis.numFmt = '#,##0'
    data = Reference(sc, min_col=1, max_col=13, min_row=7, max_row=7)
    chart.add_data(data, from_rows=True, titles_from_data=True)
    data2 = Reference(sc, min_col=1, max_col=13, min_row=9, max_row=9)
    chart.add_data(data2, from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(sc, min_col=2, max_col=13, min_row=3, max_row=3))
    chart.height, chart.width = 8, 22
    sc.add_chart(chart, "B15")

    line = LineChart()
    line.title = "RevPAR by month"
    line.add_data(Reference(sc, min_col=1, max_col=13, min_row=11, max_row=12), from_rows=True, titles_from_data=True)
    line.set_categories(Reference(sc, min_col=2, max_col=13, min_row=3, max_row=3))
    line.height, line.width = 8, 22
    sc.add_chart(line, "B32")


def _bridge_sheet(wb: Workbook) -> None:
    b = wb.create_sheet("Bridge")
    b["A1"] = "Room revenue bridge: budget to forecast (formulas)"
    b["A1"].font = Font(bold=True, size=13, color=NAVY)
    b["A2"] = ("Volume effect = (forecast rooms - budget rooms) x budget ADR; rate effect = (forecast ADR - budget ADR)"
               " x forecast rooms. The two effects reconcile exactly to the total variance.")
    for j, h in enumerate(["Month", "Budget revenue", "Volume effect", "Rate effect", "Forecast revenue",
                           "Check"]):
        b.cell(row=4, column=1 + j, value=h).font = Font(bold=True)
    for m in range(12):
        r = 5 + m
        col = get_column_letter(2 + m)
        b.cell(row=r, column=1, value=MONTHS[m])
        b.cell(row=r, column=2, value=f"=Scenarios!{col}7")
        b.cell(row=r, column=3, value=f"=(Scenarios!{col}8-Scenarios!{col}6)*'Scenario inputs'!{col}7")
        b.cell(row=r, column=4, value=f"=('Scenario inputs'!{col}9-'Scenario inputs'!{col}7)*Scenarios!{col}8")
        b.cell(row=r, column=5, value=f"=Scenarios!{col}9")
        b.cell(row=r, column=6, value=f"=ABS(B{r}+C{r}+D{r}-E{r})<0.01")
        for c in range(2, 6):
            _style_cell(b.cell(row=r, column=c), "calc", "currency0", True)
        _style_cell(b.cell(row=r, column=6), "check", "general", True)
    r = 17
    b.cell(row=r, column=1, value="Full year").font = Font(bold=True)
    for c, colL in zip(range(2, 6), "BCDE"):
        b.cell(row=r, column=c, value=f"=SUM({colL}5:{colL}16)")
        _style_cell(b.cell(row=r, column=c), "output", "currency0", True)
    b.cell(row=r, column=6, value="=ABS(B17+C17+D17-E17)<0.01")
    _style_cell(b.cell(row=r, column=6), "check", "general", True)
    b.column_dimensions["A"].width = 14
    for c in "BCDEF":
        b.column_dimensions[c].width = 18


def _checks_sheet(wb: Workbook, forecast: HotelAssumptions) -> None:
    ws = wb.create_sheet("Checks & notes")
    ws["A1"] = "Validation results at export time (from FinPilot's independent model)"
    ws["A1"].font = Font(bold=True, size=13, color=NAVY)
    for j, h in enumerate(["Check", "Status", "Category", "Detail"]):
        ws.cell(row=3, column=1 + j, value=h).font = Font(bold=True)
    for i, chk in enumerate(validate(forecast), start=4):
        ws.cell(row=i, column=1, value=chk.name)
        st = ws.cell(row=i, column=2, value=chk.status.upper())
        st.font = Font(bold=True, color={"pass": "1B7F5A", "warn": "A66A00", "fail": "C0392B"}[chk.status])
        ws.cell(row=i, column=3, value=chk.category)
        ws.cell(row=i, column=4, value=chk.message).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 90


def export_lab_sheet(path: Path, sheet: Sheet, title: str = "FinPilot Lab model") -> Path:
    wb = Workbook()
    _cover(wb, title, "Exported from the FinPilot Modeling Lab",
           ["Formulas are exported as Excel formulas and recalculate when opened (FinPilot does not run Excel).",
            "Blue cells are inputs; black cells are calculations."])
    ws = wb.create_sheet("Model")
    write_sheet(ws, sheet)
    for name, target in sheet.names.items():
        parts = target.split(":")
        ref = ":".join("$" + re.sub(r"(\d+)", r"$\1", p) for p in parts)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Model!{ref}")
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)
    return Path(path)


def read_back_sheet(path: Path, sheet_name: str = "Model") -> Sheet:
    """Load an exported worksheet back into a FinPilot Sheet (used to verify exports)."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[sheet_name]
    s = Sheet(sheet_name, rows=ws.max_row, cols=ws.max_column)
    items = []
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            v = c.value
            raw = v.replace("_xlfn.", "") if isinstance(v, str) and v.startswith("=") else (
                repr(v) if isinstance(v, float) else str(v))
            items.append(((c.row - 1, c.column - 1), raw))
    s.set_many(items, record=False)
    for name, dn in wb.defined_names.items():
        target = dn.attr_text.split("!")[-1].replace("$", "")
        s.names[name.upper()] = target
    s.recalculate()
    return s
